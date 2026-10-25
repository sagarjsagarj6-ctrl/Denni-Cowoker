"""Securely convert user text and common files to EvidenceItem records."""

from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
from pathlib import Path
from typing import Any
from uuid import uuid4

from dynamic_agents.config.settings import Settings
from dynamic_agents.core.errors import InputSecurityError
from dynamic_agents.models.domain import EvidenceItem, InputEnvelope, InputKind, NormalizedInput


class InputNormalizer:
    """Normalize supported files while keeping source/provenance metadata."""

    SUPPORTED = {
        ".txt",
        ".md",
        ".json",
        ".csv",
        ".pdf",
        ".docx",
        ".xlsx",
        ".xlsm",
        ".png",
        ".jpg",
        ".jpeg",
        ".gif",
        ".webp",
        ".wav",
        ".mp3",
        ".mp4",
        ".mov",
    }

    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def normalize(self, envelope: InputEnvelope) -> NormalizedInput:
        evidence: list[EvidenceItem] = []
        for input_file in envelope.files:
            evidence.append(self._normalize_file(Path(input_file.path), input_file.display_name))
        return NormalizedInput(
            request_id=envelope.request_id,
            text=envelope.text.strip(),
            evidence=evidence,
            company_id=envelope.company_id,
            case_id=envelope.case_id,
            metadata=envelope.metadata,
        )

    def _normalize_file(self, path: Path, display_name: str | None) -> EvidenceItem:
        resolved = path.resolve()
        if not resolved.exists() or not resolved.is_file():
            raise InputSecurityError(f"Input file does not exist: {path.name}")
        if resolved.stat().st_size > self.settings.max_file_size_bytes:
            raise InputSecurityError(f"Input file exceeds MAX_FILE_SIZE_BYTES: {resolved.name}")
        suffix = resolved.suffix.lower()
        if suffix not in self.SUPPORTED:
            raise InputSecurityError(f"Unsupported input type: {suffix or 'unknown'}")
        raw = resolved.read_bytes()
        checksum = hashlib.sha256(raw).hexdigest()
        mime = mimetypes.guess_type(resolved.name)[0] or "application/octet-stream"
        extracted_text, metadata, kind = self._extract(resolved, raw, mime)
        return EvidenceItem(
            evidence_id=f"ev_{uuid4().hex[:12]}",
            source_name=display_name or resolved.name,
            mime_type=mime,
            checksum_sha256=checksum,
            source_location=str(resolved),
            kind=kind,
            extracted_text=extracted_text,
            metadata=metadata,
            reference=resolved.name,
            trusted=False,
        )

    def _extract(self, path: Path, raw: bytes, mime: str) -> tuple[str, dict[str, Any], InputKind]:
        suffix = path.suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
            return "", {"bytes": len(raw), "vision_required": True}, InputKind.image
        if suffix in {".wav", ".mp3"}:
            return "", {"bytes": len(raw), "transcription_available": False}, InputKind.audio
        if suffix in {".mp4", ".mov"}:
            return "", {"bytes": len(raw), "frame_extraction_available": False}, InputKind.video
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader  # type: ignore[import-not-found]

                reader = PdfReader(str(path))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                return text, {"pages": len(reader.pages)}, InputKind.file
            except ImportError:
                return "", {"pages": None, "optional_dependency": "pypdf"}, InputKind.file
            except Exception as exc:
                return "", {"extraction_error": type(exc).__name__}, InputKind.file
        if suffix == ".docx":
            try:
                from docx import Document  # type: ignore[import-not-found]

                document = Document(str(path))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                return text, {"paragraphs": len(document.paragraphs)}, InputKind.file
            except ImportError:
                return "", {"optional_dependency": "python-docx"}, InputKind.file
            except Exception as exc:
                return "", {"extraction_error": type(exc).__name__}, InputKind.file
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook  # type: ignore[import-untyped]

                workbook = load_workbook(path, read_only=True, data_only=True)
                parts: list[str] = []
                rows = 0
                for sheet in workbook.worksheets:
                    parts.append(f"[sheet:{sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        rows += 1
                        parts.append(
                            " | ".join("" if value is None else str(value) for value in row)
                        )
                return (
                    "\n".join(parts),
                    {"sheets": workbook.sheetnames, "rows": rows},
                    InputKind.file,
                )
            except ImportError:
                return "", {"optional_dependency": "openpyxl"}, InputKind.file
            except Exception as exc:
                return "", {"extraction_error": type(exc).__name__}, InputKind.file
        text = raw.decode("utf-8", errors="replace")
        if suffix == ".json":
            try:
                parsed = json.loads(text)
                return (
                    json.dumps(parsed, indent=2, ensure_ascii=False),
                    {"json": True},
                    InputKind.file,
                )
            except json.JSONDecodeError:
                return text, {"json": False, "parse_error": "JSONDecodeError"}, InputKind.file
        if suffix == ".csv":
            csv_rows = list(csv.reader(text.splitlines()))
            return (
                text,
                {"columns": len(csv_rows[0]) if csv_rows else 0, "rows": len(csv_rows)},
                InputKind.file,
            )
        return text, {"bytes": len(raw)}, InputKind.file
